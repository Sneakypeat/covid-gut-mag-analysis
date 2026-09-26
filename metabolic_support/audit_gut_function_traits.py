#!/usr/bin/env python3
"""Independently read the workbook with openpyxl and recompute trait calls."""
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'result2/n76/metabolic_support/gut_traits_v1'


def table(path, delimiter='\t'):
    with path.open(newline='') as stream:
        return list(csv.DictReader(stream, delimiter=delimiter))


def state(value):
    return 'unknown' if value is None else ('positive' if value else 'evidence_not_detected')


def and3(values):
    return False if any(x is False for x in values) else (None if any(x is None for x in values) else True)


def or3(values):
    return True if any(x is True for x in values) else (None if any(x is None for x in values) else False)


def main():
    gate = json.loads((OUT/'analysis_complete.json').read_text())
    assert gate['complete'] and gate['primary_n']==294 and gate['trait_n']==5
    for key,path in gate['input_paths'].items():
        assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==gate['input_sha256'][key], key
    for name,digest in gate['output_sha256'].items():
        assert hashlib.sha256((OUT/name).read_bytes()).hexdigest()==digest, name
    actual = table(OUT/'genome_traits_wide.tsv')
    genomes = {row['catalog_id'] for row in actual}
    assert len(actual)==len(genomes)==584
    terms = ('Amorphous Cellulose','Crystalline Cellulose','Arabinan',
             'Beta-galactan (pectic galactan)','Beta-mannan','Starch','Pectin',
             'Xylan','Xyloglucan','Mixed-Linkage glucans')
    needed = {'K00634','K00929','K00625','K00925','GH33','GH29','GH95','GH20','GH35'}
    workbook = load_workbook(gate['input_paths']['workbook'], read_only=True, data_only=True)
    counts = {}; plant = set()
    for name in ['MISC','carbon utilization']:
        rows = workbook[name].iter_rows(values_only=True)
        header = next(rows)
        positions = {g:header.index(g) for g in genomes}
        gene_col = header.index('gene_id'); sub_col = header.index('subheader')
        for row in rows:
            gene = row[gene_col]
            selected = any(f'{term} Backbone Cleavage' in (row[sub_col] or '') for term in terms)
            if selected: plant.add(gene)
            if gene not in needed and not selected: continue
            vector = {g:None if row[i] in {None,''} else float(row[i]) for g,i in positions.items()}
            if gene in counts: assert counts[gene]==vector,gene
            counts[gene]=vector
    workbook.close()
    assert plant==set(gate['plant_marker_ids'])
    support=json.loads(Path(gate['input_paths']['sets']).read_text())
    primary={r['catalog_id'] for r in table(Path(gate['input_paths']['primary']))}
    group_counts=Counter(); recalculated=0
    for row in actual:
        g=row['catalog_id']
        positive=lambda gene: None if counts[gene][g] is None else counts[gene][g]>0
        expected={
            'plant_backbone_cazyme':or3([positive(x) for x in plant]),
            'butyrate_terminal_pair':and3([positive('K00634'),positive('K00929')]),
            'acetate_terminal_pair':and3([positive('K00625'),positive('K00925')]),
            'mucin_gh_repertoire':and3([positive('GH33'),or3([positive('GH29'),positive('GH95')]),or3([positive('GH20'),positive('GH35')])]),
            'strict_acid_uptake':None if g not in support else bool(set(support[g]['requirements']) & {'lac__D','lac__L','succ','ac'})}
        for trait,value in expected.items():
            assert row[trait]==state(value),(g,trait)
            recalculated+=1
            if g in primary:
                group_counts[(trait,row['ancom_direction'],'n_group')]+=1
                group_counts[(trait,row['ancom_direction'],'n_positive')]+=value is True
                group_counts[(trait,row['ancom_direction'],'n_unknown')]+=value is None
    for row in table(OUT/'trait_group_summary.tsv'):
        for field in ['n_group','n_positive','n_unknown']:
            assert int(row[field])==group_counts[(row['trait_id'],row['ancom_direction'],field)]
        assert abs(float(row['fraction_positive_all'])-int(row['n_positive'])/int(row['n_group']))<1e-12
    print(json.dumps({'status':'pass','catalogue_n':584,'primary_n':294,'traits':5,
                      'trait_calls_independently_recomputed':recalculated,
                      'workbook_parser':'openpyxl, independent of builder XML parser',
                      'output_hashes_verified':len(gate['output_sha256']),
                      'scope':'Annotation arithmetic and provenance; not validation of ecological guild membership'},indent=2))


if __name__=='__main__':
    main()
